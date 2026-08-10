import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    """Establece el color de fondo de una celda en una tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Ajusta los márgenes internos (padding) de una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def construir_documento_word():
    doc = docx.Document()

    # Configuración de márgenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilos de fuentes
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # --- TÍTULO PRINCIPAL ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Documentación Técnica Completa\nSistema de Gestión de Estacionamiento")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78) # Azul marino
    p_title.paragraph_format.space_after = Pt(24)

    # --- SECCIÓN 1 ---
    h1 = doc.add_heading("1. Resumen Ejecutivo del Proyecto", level=1)
    h1.style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    
    p = doc.add_paragraph()
    p.add_run("El ").font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    r_bold = p.add_run("Sistema de Gestión de Estacionamiento")
    r_bold.bold = True
    p.add_run(" es una aplicación de escritorio multiplataforma desarrollada en ")
    p.add_run("Python").bold = True
    p.add_run(", utilizando ")
    p.add_run("Flet").bold = True
    p.add_run(" (basado en Flutter) para la interfaz gráfica de usuario y ")
    p.add_run("SQLite").bold = True
    p.add_run(" como motor de base de datos relacional.\n\n"
              "La aplicación permite la autenticación y registro de usuarios (conductores y administradores), "
              "la gestión de perfiles, la visualización de lugares de estacionamiento organizados por sectores o zonas, "
              "el registro de vehículos, y la gestión transaccional de reservas y pagos.")

    # --- SECCIÓN 2 ---
    h2 = doc.add_heading("2. Arquitectura General y Estructura del Proyecto", level=1)
    h2.style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    
    doc.add_paragraph(
        "El sistema adopta una arquitectura modular en capas inspirada en el patrón DAO (Data Access Object) / MVC (Modelo-Vista-Controlador). "
        "Esta estructura desacopla completamente la lógica visual (UI) de la lógica de persistencia y base de datos."
    )

    doc.add_heading("Estructura de Directorios del Proyecto", level=2)
    
    directorios = [
        ("Aplicacion_De_Gestion/", "Directorio raíz del proyecto."),
        ("├── main.py", "Capa de Presentación (UI): Interfaz Flet, navegación SPA y controladores."),
        ("├── requirements.txt", "Dependencias requeridas (flet, etc.)."),
        ("└── database/", "Paquete de Base de Datos y Lógica de Negocio."),
        ("    ├── conexion.py", "Infraestructura: Configuración de SQLite, Pragma de claves foráneas y DDL."),
        ("    ├── consultas.py", "Lógica de Negocio: Operaciones DML parametrizadas para usuarios, plazas y pagos."),
        ("    └── app_gestion.db", "Base de datos relacional SQLite autogenerada.")
    ]

    for item, desc in directorios:
        p_dir = doc.add_paragraph()
        p_dir.paragraph_format.left_indent = Inches(0.2)
        p_dir.paragraph_format.space_after = Pt(2)
        r_item = p_dir.add_run(f"{item:<25} ")
        r_item.font.name = 'Consolas'
        r_item.font.size = Pt(9.5)
        r_item.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
        p_dir.add_run(f"— {desc}")

    # --- SECCIÓN 3 ---
    h3 = doc.add_heading("3. Esquema y Diseño de la Base de Datos", level=1)
    h3.style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    doc.add_paragraph(
        "La base de datos relacional consta de 10 tablas relacionales diseñadas para garantizar la integridad referencial "
        "mediante el uso explícito de FOREIGN KEY y restricciones de unicidad (UNIQUE)."
    )

    doc.add_heading("Tabla de Relaciones e Integridad Referencial", level=2)

    # Tabla en Word: Relaciones
    rel_data = [
        ("Tabla Origen", "Card.", "Tabla Destino", "Clave Foránea (FK)", "Propósito"),
        ("TipoUsuario", "1 : N", "Usuarios", "Usuarios.id_rol", "Asigna el rol de acceso (Admin/Conductor)."),
        ("Usuarios", "1 : 1", "Conductores", "Conductores.id_usuario", "Vincula credenciales con el perfil del conductor."),
        ("Usuarios", "1 : 1", "Administradores", "Administradores.id_usuario", "Vincula credenciales con el perfil administrativo."),
        ("Conductores", "1 : N", "Vehiculo", "Vehiculo.id_conductor", "Relaciona los vehículos con su propietario."),
        ("TipoVehiculo", "1 : N", "Vehiculo", "Vehiculo.id_tipo_vehiculo", "Clasifica el vehículo (Auto, Moto, Camioneta)."),
        ("TipoVehiculo", "1 : N", "lugares", "lugares.id_tipo_vehiculo", "Restringe tipo de vehículo permitido por plaza."),
        ("EstadoLugar", "1 : N", "lugares", "lugares.id_estado", "Define si la plaza está Libre, Ocupada o Reservada."),
        ("Conductores", "1 : N", "Reservas", "Reservas.id_conductor", "Registra el autor de la reserva."),
        ("lugares", "1 : N", "Reservas", "Reservas.id_lugar", "Identifica la plaza de estacionamiento reservada."),
        ("Vehiculo", "1 : N", "Reservas", "Reservas.id_vehiculo", "Especifica la patente asociada a la reserva."),
        ("Reservas", "1 : N", "Pagos", "Pagos.id_reserva", "Guarda la transacción financiera de la reserva.")
    ]

    t_rel = doc.add_table(rows=len(rel_data), cols=5)
    t_rel.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row in enumerate(t_rel.rows):
        for j, cell in enumerate(row.cells):
            cell.text = rel_data[i][j]
            set_cell_margins(cell)
            if i == 0:
                set_cell_background(cell, "1F4E78") # Encabezado azul
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                bg = "F2F4F7" if i % 2 == 0 else "FFFFFF"
                set_cell_background(cell, bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Especificación de Tablas
    doc.add_heading("Especificación de Tablas y Columnas", level=2)

    tablas_spec = [
        ("A. Tablas Catálogo / Parámetros", [
            ("1. TipoUsuario", "id_rol (PK), rol_usuario (UNIQUE, NOT NULL: 'Administrador', 'Conductor')"),
            ("2. EstadoLugar", "id_estado (PK), descripcion (UNIQUE, NOT NULL: 'Disponible', 'Ocupado', 'Reservado')"),
            ("3. TipoVehiculo", "id_tipo_vehiculo (PK), tipo_vehiculo (UNIQUE, NOT NULL: 'Auto', 'Moto', 'Camioneta')")
        ]),
        ("B. Tablas de Usuarios y Entidades", [
            ("4. Usuarios", "id_usuario (PK), mail (UNIQUE), contraseña, id_rol (FK -> TipoUsuario)"),
            ("5. Conductores", "id_conductor (PK), nombre, dni (UNIQUE), id_usuario (UNIQUE, FK -> Usuarios)"),
            ("6. Administradores", "id_administrador (PK), nombre, id_usuario (UNIQUE, FK -> Usuarios)")
        ]),
        ("C. Tablas Operativas y de Dominio", [
            ("7. lugares", "id_lugar (PK), num_lugar, piso_o_sector, precio, id_tipo_vehiculo (FK), id_estado (FK)"),
            ("8. Vehiculo", "id_vehiculo (PK), patente (UNIQUE), id_conductor (FK), id_tipo_vehiculo (FK)"),
            ("9. Reservas", "id_reserva (PK), estado_reserva, fecha_reserva, hora_inicio, hora_fin, id_conductor (FK), id_lugar (FK), id_vehiculo (FK)"),
            ("10. Pagos", "id_pago (PK), monto, metodo_pago, fecha_pago, id_reserva (FK -> Reservas)")
        ])
    ]

    for cat_title, items in tablas_spec:
        h_cat = doc.add_heading(cat_title, level=3)
        h_cat.style.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        for t_name, t_cols in items:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Inches(0.2)
            r_name = p_item.add_run(f"• {t_name}: ")
            r_name.bold = True
            p_item.add_run(t_cols)

    # --- SECCIÓN 4 ---
    h4 = doc.add_heading("4. Desglose Módulo por Módulo y Flujo de Datos", level=1)
    h4.style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    doc.add_heading("A. Flujo de Navegación del Usuario (SPA)", level=2)
    
    pasos_flujo = [
        ("Paso 1: Inicio de Sesión", "El usuario ingresa correo y contraseña. Se consulta verificar_credenciales() en la DB. Si es correcto, pasa al Menú Principal."),
        ("Paso 2: Registro de Conductor", "Si no tiene cuenta, completa Nombre, DNI, Correo y Contraseña. Se guardan en las tablas Usuarios y Conductores de forma relacional."),
        ("Paso 3: Menú Principal", "Muestra un saludo personalizado con los datos del usuario en sesión y ofrece opciones principales (Buscar Estacionamiento, Cerrar Sesión)."),
        ("Paso 4: Búsqueda y Reserva", "El conductor explora plazas por Zona/Sector, elige un lugar libre y registra la reserva junto con el método de pago.")
    ]

    for paso, desc in pasos_flujo:
        p_paso = doc.add_paragraph()
        p_paso.paragraph_format.left_indent = Inches(0.2)
        r_p = p_paso.add_run(f"► {paso}: ")
        r_p.bold = True
        r_p.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        p_paso.add_run(desc)

    doc.add_heading("B. Módulo de Infraestructura: conexion.py", level=2)
    doc.add_paragraph(
        "Responsable de conectar a la base de datos app_gestion.db, activar explícitamente "
        "el soporte de claves foráneas (PRAGMA foreign_keys = ON) y ejecutar la creación idempotente "
        "de tablas con INSERT OR IGNORE para poblar los catálogos iniciales."
    )

    doc.add_heading("C. Módulo de Lógica de Negocio: consultas.py", level=2)
    doc.add_paragraph(
        "Centraliza todas las consultas SQL preparadas (parametrizadas con ?) para prevenir inyección SQL. "
        "Implementa transacciones con try / commit / rollback para garantizar que la falla en una sub-inserción "
        "(ej. DNI duplicado) no deje registros huérfanos."
    )

    # --- SECCIÓN 5 ---
    h5 = doc.add_heading("5. Guía de Instalación y Ejecución", level=1)
    h5.style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    pasos_inst = [
        ("1. Instalar Flet", "pip install flet"),
        ("2. Inicializar Base de Datos", "py database/conexion.py"),
        ("3. Ejecutar la Aplicación", "py main.py")
    ]

    for paso, cmd in pasos_inst:
        p_i = doc.add_paragraph()
        p_i.paragraph_format.left_indent = Inches(0.2)
        p_i.add_run(f"• {paso}: ").bold = True
        r_cmd = p_i.add_run(cmd)
        r_cmd.font.name = 'Consolas'
        r_cmd.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)

    # --- SECCIÓN 6 ---
    h6 = doc.add_heading("6. Diagnóstico y Solución de Problemas Frecuentes", level=1)
    h6.style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    diag_data = [
        ("Error / Síntoma", "Causa Raíz", "Solución Aplicada"),
        ("FOREIGN KEY constraint failed", "Las tablas catálogo (ej. TipoUsuario) estaban vacías al vincular un registro.", "Se aplicó Defensive Database Pattern con INSERT OR IGNORE de los roles previos."),
        ("ImportError: cannot import name ...", "Cambios no guardados en el archivo o caché obsoleta de Python.", "Guardar con Ctrl+S y eliminar la carpeta database/__pycache__."),
        ("No se encontró Python en Windows", "Alias de ejecución activo apuntando a Microsoft Store.", "Usar 'py main.py' o desactivar el alias en la Configuración de Windows.")
    ]

    t_diag = doc.add_table(rows=len(diag_data), cols=3)
    t_diag.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row in enumerate(t_diag.rows):
        for j, cell in enumerate(row.cells):
            cell.text = diag_data[i][j]
            set_cell_margins(cell)
            if i == 0:
                set_cell_background(cell, "1F4E78")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                bg = "F2F4F7" if i % 2 == 0 else "FFFFFF"
                set_cell_background(cell, bg)

    # Guardar documento
    nombre_archivo = "Documentacion_Sistema_Estacionamiento.docx"
    doc.save(nombre_archivo)
    print(f"¡Documento creado exitosamente como '{nombre_archivo}'!")

if __name__ == "__main__":
    construir_documento_word()